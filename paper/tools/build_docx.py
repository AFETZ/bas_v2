#!/usr/bin/env python3
"""Create an anonymous macro-free DOCX using the official template's styles.

The original DOTM is read as OOXML data; its body, metadata, headers, objects,
relationships and macros are never imported or executed. Explicit typography
follows the template's own prose requirements where inherited size is absent.
"""
from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
from datetime import datetime, timezone
import csv
import json
import re
import hashlib
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BUILD=ROOT/'build'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
TABLE_SEQUENCE=[]
TABLE_EN_CAPTIONS={
    'related':'Scope of related work; comparison does not establish absence of unlisted capabilities.',
    'six_cells':'Main 30-run study, five runs per cell. Means; times start at actual RF fault. Dashes mean not applicable.',
    'decomposition':'Mean T_goal for the actual last-finishing UAV, five runs per policy; first UART differs from the latest completion refresh.',
    'sensitivity':'Separate nominal-outage study, three planned runs per cell. Observed UART means include n; censored or missing outcomes are not zero.',
    'operator':'Separate console experiment: complete parameter coverage for all five vehicles; only UAV3 flew.'}

def style(doc,name,font,size,bold=False,italic=False):
    s=doc.styles[name]
    s.font.name=font;s.font.size=Pt(size);s.font.bold=bold;s.font.italic=italic
    s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font)
    return s

def configure(doc):
    with ZipFile(ROOT/'template/official_russian_template.dotm') as z:
        official_styles=etree.fromstring(z.read('word/styles.xml'))
        section=etree.fromstring(z.read('word/document.xml')).find('.//{%s}sectPr'%W)
    # A fresh package avoids all sample authors, hidden objects and linked paths.
    current=doc.styles.element
    for child in list(current):current.remove(child)
    for child in official_styles:current.append(deepcopy(child))
    sec=doc.sections[0]
    for tag in ('pgSz','pgMar','cols'):
        existing=sec._sectPr.find(qn('w:'+tag))
        if existing is not None:sec._sectPr.remove(existing)
        node=section.find(qn('w:'+tag))
        if node is not None:sec._sectPr.append(deepcopy(node))
    for node in sec._sectPr.findall(qn('w:pgNumType')):sec._sectPr.remove(node)
    pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'1');sec._sectPr.append(pg)
    normal=style(doc,'Normal','Times New Roman',10)
    normal.paragraph_format.line_spacing=1
    main=style(doc,'ispText_main','Times New Roman',10)
    main.paragraph_format.line_spacing=1
    main.paragraph_format.space_before=Pt(1.5);main.paragraph_format.space_after=Pt(1.5)
    main.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    main.paragraph_format.first_line_indent=Cm(0)
    main.paragraph_format.widow_control=True
    title=style(doc,'ispHeader','Arial',16,True)
    title.paragraph_format.space_before=Pt(48);title.paragraph_format.space_after=Pt(24)
    title.paragraph_format.page_break_before=False
    title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    eng=style(doc,'ispHeader1','Arial',16,True)
    eng.paragraph_format.space_before=Pt(18);eng.paragraph_format.space_after=Pt(12)
    eng.paragraph_format.page_break_before=False
    eng.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for name,italic in [('ispSubHeader-2 level',True),('ispSubHeader-3 level',False)]:
        s=style(doc,name,'Arial',11,True,italic)
        s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(3)
        s.paragraph_format.left_indent=Cm(0);s.paragraph_format.first_line_indent=Cm(0)
        s.paragraph_format.keep_with_next=True;s.paragraph_format.page_break_before=False
        s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        np=s._element.get_or_add_pPr().find(qn('w:numPr'))
        if np is not None:np.getparent().remove(np)
        # Suppress inherited automatic numbering; numbers are source text.
        np=OxmlElement('w:numPr');n=OxmlElement('w:numId');n.set(qn('w:val'),'0');np.append(n)
        s._element.get_or_add_pPr().append(np)
    s=style(doc,'ispAnotation','Times New Roman',9)
    s.paragraph_format.space_before=Pt(6);s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1
    s.paragraph_format.widow_control=True
    style(doc,'ispAuthor','Times New Roman',10,italic=True)
    style(doc,'ispPicture_sign','Times New Roman',9,italic=True)
    style(doc,'ispLitList','Times New Roman',9)
    # An anonymous running title and a live page field replace all sample data.
    p=sec.header.paragraphs[0];p.text='Восстановление управления через бортовой ретранслятор'
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(8);r.font.name='Times New Roman'
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    cp=doc.core_properties
    cp.author='';cp.last_modified_by='';cp.title='Anonymous UAV relay manuscript'
    cp.subject='';cp.comments='';cp.keywords='';cp.category='';cp.identifier=''
    cp.created=datetime(2026,10,27,tzinfo=timezone.utc)
    cp.modified=datetime(2026,10,27,tzinfo=timezone.utc)
    cp.revision=1

def p(doc,text,sty='ispText_main'):
    para=doc.add_paragraph(text,sty)
    if sty in ('ispHeader','ispHeader1'):
        np=OxmlElement('w:numPr');n=OxmlElement('w:numId');n.set(qn('w:val'),'0');np.append(n)
        para._p.get_or_add_pPr().append(np)
    return para

def figure_caption(doc,text,english=False):
    para=p(doc,text,'ispPicture_sign')
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=para.paragraph_format
    pf.space_before=Pt(0 if english else 6)
    pf.space_after=Pt(6 if english else 0)
    pf.keep_with_next=not english
    return para

def table(doc,name):
    TABLE_SEQUENCE.append(name)
    if name=='related':
        caption='Табл. 1. Сопоставление фокуса работ; наличие возможностей не опровергается по их отсутствию в кратком описании'
        heads=['Работа','Основной подход','Отношение к данной работе']
        rows=[['FlyNetSim [1]','ArduPilot + ns-3, middleware, ad-hoc','Прямое предшествующее решение интеграции'],
              ['ROS-NetSim [2]','ROS-интерфейс робототехники и сети','Разделение моделей и PAC-цикл'],
              ['Flightmare [3]','Отдельные rendering и dynamics','Компромисс скорости и сложности'],
              ['Данная работа','SITL → native radio → final UART gate','Парный отказный опыт; команда и физическая цель раздельно']]
    elif name=='six_cells':
        caption='Табл. 2. Основная серия: по пять запусков в ячейке. Средние; T отсчитывается от фактического отказа. «—» означает неприменимость без отказа'
        heads=['Отказ','Политика','Timely / offered','T_UART, с','T_goal, с']
        with (ROOT/'data/six_cells.csv').open(newline='') as f:source=list(csv.DictReader(f))
        rows=[]
        for r in source:
            fmt=lambda k: f"{float(r[k]):.3f}".replace('.',',') if r[k] else '—'
            rows.append([r['condition'],r['policy'],f"{100*float(r['mean_timely_per_offered']):.3f}%".replace('.',','),
                        fmt('mean_actual_control_s'),fmt('mean_actual_goal_s')])
    elif name=='decomposition':
        caption='Табл. 3. Среднее T_goal по последнему завершившему UAV (n=5); первый UART той же цели, не последний refresh CID'
        heads=['Политика','До создания, с','До первого UART, с','Движение + hold, с','Сумма, с']
        with (ROOT/'data/decomposition.csv').open(newline='') as f:source=list(csv.DictReader(f))
        keys=['fault_to_goal_creation_s','goal_creation_to_first_UART_s','first_UART_to_completion_s','fault_to_completion_s']
        rows=[]
        for policy in ('fixed','failover'):
            selected=[r for r in source if r['policy']==policy and r['group']=='all_five']
            assert len(selected)==5
            rows.append([policy]+[f"{sum(float(r[k]) for r in selected)/5:.3f}".replace('.',',') for k in keys])
    elif name=='sensitivity':
        caption='Табл. 4. Отдельная серия чувствительности; по три запланированных запуска. T_UART — среднее наблюдений с их числом n; пропуски и цензурирование не равны нулю'
        heads=['Номинал, с','Политика','Actual RF min–max, с','Timely, %','T_UART, с (n)','Censored / failed']
        with (ROOT/'data/sensitivity_runs.csv').open(newline='') as f:source=list(csv.DictReader(f))
        assert len(source)==18
        rows=[]
        for duration in (.5,2,10):
            for policy in ('fixed','failover'):
                selected=[r for r in source if float(r['outage_duration_s'])==duration and r['policy']==policy]
                assert len(selected)==3
                timely=[float(r['timely_per_offered'])*100 for r in selected if r['timely_per_offered']]
                uart=[float(r['common_UART_recovery_s']) for r in selected
                      if r['common_UART_status']=='observed' and r['common_UART_recovery_s']]
                cen=sum(r['common_UART_status']!='observed' and bool(r['common_UART_censor_time_s']) for r in selected)
                failed=sum(r['run_status']!='completed' for r in selected)
                rf=[float(r['actual_RF_duration_s']) for r in selected if r['actual_RF_duration_s']]
                fmt=lambda vals:f"{sum(vals)/len(vals):.3f}".replace('.',',') if vals else '—'
                rf_range=f'{min(rf):.6f}–{max(rf):.6f}'.replace('.',',') if rf else '—'
                rows.append([str(duration).replace('.',','),policy,
                    rf_range,f'{fmt(timely)} ({len(timely)})',f'{fmt(uart)} ({len(uart)})',f'{cen} / {failed}'])
    elif name=='operator':
        caption='Табл. 5. Отдельный операторский опыт: полные индексированные снимки всех пяти, физический полёт только UAV3'
        heads=['Проверка','Фактический результат']
        source=json.loads((ROOT/'data/operator03_summary.json').read_text())
        before=source['parameters']['before'];after=source['parameters']['after']
        fmt=lambda x:f'{x:.3f}'.replace('.',',')
        rows=[['Версия / режим','568e8ef / headless; UAV3, остальные на земле'],
              ['Параметры до',f"5 × 1388 индексов, 0 missing; {fmt(before['duration_wall_s'])} внешних с"],
              ['Параметры после',f"5 × 1388 индексов, 0 missing; {fmt(after['duration_wall_s'])} внешних с"],
              ['Изменения / запись',f"{source['before_after_changes']} совпавших изменений; 0 PARAM_SET"],
              ['Командная цепочка','6/6 одинаковых SHA256 на всех пяти этапах; UART внутри lease'],
              ['Полёт / удержание','16→20 м relative-alt, BRAKE 7,1 с, LAND/disarm; позы отдельно']]
    else:raise ValueError(name)
    caption=re.sub(r'^Табл\. \d+\.',f'Табл. {len(TABLE_SEQUENCE)}.',caption).rstrip('.')+'.'
    caption_en=f'Table {len(TABLE_SEQUENCE)}. '+TABLE_EN_CAPTIONS[name]
    for english,text in [(False,caption),(True,caption_en)]:
        para=p(doc,text,'ispPicture_sign')
        para.alignment=WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.keep_with_next=True
        para.paragraph_format.space_before=Pt(0 if english else 6)
        para.paragraph_format.space_after=Pt(6 if english else 0)
    t=doc.add_table(rows=1,cols=len(heads));t.style='Table Grid'
    t.autofit=False
    for c,txt in zip(t.rows[0].cells,heads):c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=txt
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(2);para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1
                for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(9);run.bold=i==0
    widths=([3.9,10.0] if name=='operator' else [2.6,5.0,6.3] if name=='related'
            else [1.5,1.6,3.0,2.2,3.2,2.4] if name=='sensitivity' else [2.5,2.2,3.3,2.9,3.0])
    assert sum(widths)<=14
    for index,w in enumerate(widths):
        t.columns[index].width=Cm(w)
        for row in t.rows:row.cells[index].width=Cm(w)

def build():
    TABLE_SEQUENCE.clear()
    BUILD.mkdir(exist_ok=True)
    doc=Document();configure(doc)
    for line in (ROOT/'src/manuscript.md').read_text().splitlines():
        if not line.strip():continue
        if line.startswith('@title '):p(doc,line[7:],'ispHeader')
        elif line.startswith('@authors '):p(doc,line[9:],'ispAuthor')
        elif line.startswith('@abstract '):p(doc,'Аннотация. '+line[10:],'ispAnotation')
        elif line.startswith('@keywords '):p(doc,'Ключевые слова: '+line[10:],'ispAnotation')
        elif line.startswith('@title_en '):p(doc,line[10:],'ispHeader1')
        elif line.startswith('@authors_en '):p(doc,line[12:],'ispAuthor')
        elif line.startswith('@abstract_en '):p(doc,'Abstract. '+line[13:],'ispAnotation')
        elif line.startswith('@keywords_en '):p(doc,'Keywords: '+line[13:],'ispAnotation')
        elif line.startswith('## '):p(doc,line[3:],'ispSubHeader-3 level')
        elif line.startswith('# '):p(doc,line[2:],'ispSubHeader-2 level')
        elif line.startswith('@figure '):
            _,name,width=line.split()
            if float(width)>14:raise ValueError('Official template figure width cap is 14 cm')
            para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next=True
            para.add_run().add_picture(str(ROOT/'figures'/f'{name}.png'),width=Cm(float(width)))
        elif line.startswith('@caption '):figure_caption(doc,line[9:])
        elif line.startswith('@caption_en '):figure_caption(doc,line[12:],english=True)
        elif line.startswith('@table '):table(doc,line[7:])
        elif line.startswith('@pending '):
            _,key,text=line.split(' ',2)
            para=p(doc,text,'ispAnotation')
            for r in para.runs:r.font.color.rgb=RGBColor.from_string('8B4513')
        elif line=='@references':
            p(doc,'Список литературы / References','ispSubHeader-2 level')
            for ref in (ROOT/'refs/references.txt').read_text().splitlines():p(doc,ref,'ispLitList')
        elif line.startswith('@'):raise ValueError(line)
        else:p(doc,line)
    dest=BUILD/'manuscript_anonymous.docx';doc.save(dest)
    # Fresh python-docx core is anonymous; remove default company and local path data.
    from tempfile import NamedTemporaryFile
    with ZipFile(dest) as z:parts={i.filename:z.read(i.filename) for i in z.infolist()}
    app=etree.fromstring(parts['docProps/app.xml'])
    for node in app.iter():
        if etree.QName(node).localname in ('Company','Manager'):node.text=''
    parts['docProps/app.xml']=etree.tostring(app,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(dest,'w') as z:
        for name,data in parts.items():z.writestr(name,data)
    forbidden=[b'/home/',b'Ivanov',b'Petrov','Иванов'.encode(),'Петров'.encode(),b'vbaProject',b'attachedTemplate']
    leaks=[]
    for name,data in parts.items():
        if name.endswith('.xml') or name.endswith('.rels'):
            leaks += [(name,x.decode('utf-8','replace')) for x in forbidden if x in data]
    if leaks:raise RuntimeError(f'Anonymous package scan failed: {leaks}')
    (BUILD/'docx_check.json').write_text(json.dumps({'macro_free':True,'anonymous_package_scan':leaks,
        'template_style_source':'official_russian_template.dotm','main_font_pt':10,
        'page_size_cm':[16.5,24.0],'personal_template_parts_imported':False},indent=2)+'\n')
    inputs=[ROOT/'src/manuscript.md',ROOT/'refs/references.txt',ROOT/'template/official_russian_template.dotm',
            *sorted((ROOT/'figures').glob('*.png'))]
    (BUILD/'build_manifest.json').write_text(json.dumps({
        'inputs_sha256':{str(path.relative_to(ROOT)):hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs},
        'docx_sha256':hashlib.sha256(dest.read_bytes()).hexdigest()},indent=2)+'\n')
    print(dest)

if __name__=='__main__':build()
