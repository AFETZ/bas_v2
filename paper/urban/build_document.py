#!/usr/bin/env python3
"""Build a separate anonymous revision from real analysis inputs or explicit gaps.

Run inside the existing document-tools image under the BAS runtime lock. This
script reads old evidence/styles and never writes the previous paper/build.
"""
import argparse
import csv
from datetime import datetime, timezone
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
sys.path.insert(0, str(HERE))
import study_inputs
SPEC = importlib.util.spec_from_file_location("old_paper_styles", PAPER / "tools/build_docx.py")
STYLE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(STYLE)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def resolved(path, parent):
    path = Path(path)
    return path if path.is_absolute() else (parent / path).resolve()


def fmt(value, places=3):
    if value is None or value == "":
        return "—"
    number = float(value)
    if round(number, places) == 0:
        number = 0.0
    return f"{number:.{places}f}".replace(".", ",")


class Revision:
    def __init__(self, args):
        self.args = args
        self.inputs = {}
        self.base = args.manifest.resolve().parent if args.manifest else HERE
        self.manifest = self.read_json(args.manifest) if args.manifest else {"kind": "preparation"}
        if args.final and args.supplement:
            raise ValueError("supplement.md is an archived preparation source, not the final manuscript")
        if self.manifest.get("figure_manifest"):
            rendered = self.read_json(resolved(self.manifest["figure_manifest"], self.base))
            if args.final and (rendered.get("batch_complete") is not True
                               or rendered.get("preview") is not False or rendered.get("not_rendered")):
                raise ValueError("final manuscript requires the complete measured figure batch")
            self.manifest["figures"] = rendered.get("figures", [])
        if self.manifest.get("figure_ids"):
            figures = {entry["id"]: entry for entry in self.manifest.get("figures", [])}
            selected = self.manifest["figure_ids"]
            if len(set(selected)) != len(selected) or any(name not in figures for name in selected):
                raise ValueError("manuscript figure selection must name distinct measured figures")
            self.manifest["figures"] = [figures[name] for name in selected]
        for entry in self.manifest.get("figures", []):
            override = self.manifest.get("figure_captions", {}).get(entry.get("id"), {})
            for name in ("caption_ru", "caption_en"):
                if name in override:
                    entry[name] = override[name]
        for name in self.manifest.get("evidence_sources", []):
            path = resolved(name, self.base)
            self.inputs[str(path)] = sha(path)
        self.kind = self.manifest.get("kind", "preparation")
        if self.kind not in ("preparation", "engineering_pilot", "partial_series", "frozen_series"):
            raise ValueError("unknown dataset kind")
        self.freeze = self.read_json(resolved(self.manifest["freeze_path"], self.base)) if self.manifest.get("freeze_path") else {}
        self.study_path = resolved(self.manifest["study_json"], self.base) if self.manifest.get("study_json") else None
        self.study = self.read_json(self.study_path) if self.study_path else None
        self.input_root = resolved(self.manifest["input_root"], self.base) if self.manifest.get("input_root") else (
            self.study_path.parent if self.study_path else self.base)
        self.runs = []
        for index, entry in enumerate(self.manifest.get("analyses", []), 1):
            path = resolved(entry["path"], self.base)
            result = self.read_json(path)
            self.runs.append((entry.get("label", f"U{index:02d}"), entry.get("group", "unspecified"), result))
        seen = {run["run_id"] for _, _, run in self.runs}
        if self.study:
            for index, record in enumerate(self.study.get("runs", []), 1):
                path = resolved(record["input_paths"]["analysis"], self.input_root)
                if path.is_file() and record["run_id"] not in seen:
                    if sha(path) != record.get("input_sha256", {}).get("analysis"):
                        raise ValueError(f"{record['run_id']}: analysis changed since aggregate; re-export")
                    result = self.read_json(path)
                    if result.get("run_id") != record["run_id"]:
                        raise ValueError("study/per-run identity mismatch")
                    self.runs.append((f"U{index:02d}", record.get("cohort", "unspecified"), result))
                    seen.add(record["run_id"])
        if args.final:
            if not self.manifest.get("study_json") or not self.runs:
                raise ValueError("--final requires study_json and actual analysis inputs")
            freeze = resolved(self.manifest["freeze_path"], self.base)
            self.read_json(freeze)
            controllers = {}
            for record in self.study.get("runs", []):
                path = resolved(record["input_paths"]["controller"], self.input_root)
                if not path.is_file():
                    raise ValueError(f"missing final controller: {record['run_id']}")
                controllers[record["run_id"]] = self.read_json(path)
                if sha(path) != record.get("input_sha256", {}).get("controller"):
                    raise ValueError(f"{record['run_id']}: controller changed since aggregate; re-export")
            study_inputs.validate_closed_study(self.manifest, self.freeze, self.study, controllers, sha(freeze))
            for label, _group, run in self.runs:
                checks = run.get("physical_checks", {})
                full_window = run.get("observation_window", {}).get("complete") is True
                if (checks.get("total") != checks.get("verified") or (full_window and
                        (run.get("data_quality_errors") or run.get("missing_measurements")))):
                    raise ValueError(f"{label}: unresolved recording/physical-evidence issue; keep draft")
        self.doc = Document()
        STYLE.configure(self.doc)
        self.doc.sections[0].header.paragraphs[0].text = "Восстановление управления и геометрическое экранирование"
        self.doc.core_properties.title = "Anonymous UAV relay manuscript: geometry extension"
        self.doc.core_properties.modified = datetime(2026, 10, 29, tzinfo=timezone.utc)
        self.table_number, self.figure_number = 0, 0

    def read_json(self, path):
        path = Path(path).resolve()
        self.inputs[str(path)] = sha(path)
        return json.loads(path.read_text())

    def para(self, text, style="ispText_main"):
        return STYLE.p(self.doc, text, style)

    def table(self, heads, rows, caption_ru, caption_en):
        self.table_number += 1
        for text in (f"Табл. {self.table_number}. {caption_ru}", f"Table {self.table_number}. {caption_en}"):
            p = self.para(text, "ispPicture_sign")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
        table = self.doc.add_table(rows=1, cols=len(heads))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, heads):
            cell.text = str(text)
        for row in rows:
            if len(row) != len(heads):
                raise ValueError("table column count mismatch")
            for cell, text in zip(table.add_row().cells, row):
                cell.text = "—" if text is None or text == "" else str(text)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.name, run.font.size, run.bold = "Times New Roman", Pt(9), i == 0

    def figure(self, path, caption_ru, caption_en):
        path = Path(path).resolve()
        self.inputs[str(path)] = sha(path)
        from PIL import Image
        with Image.open(path) as image:
            if image.width / (13.9 / 2.54) < 300:
                raise ValueError(f"figure has less than 300 effective dpi: {path}")
        self.figure_number += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(13.9))
        STYLE.figure_caption(self.doc, f"Рис. {self.figure_number}. {caption_ru}")
        STYLE.figure_caption(self.doc, f"Fig. {self.figure_number}. {caption_en}", english=True)

    def methodology(self):
        cases = self.freeze.get("cases", [])
        cohorts = {name: sum(r.get("cohort") == name for r in cases)
                   for name in ("main", "sensitivity", "engineering_boundary")}
        blocks = {r["block_id"] for r in cases if r.get("cohort") == "main"}
        repeats = (f"{len(blocks)} блока, {cohorts['main']} основных; "
                   f"+{cohorts['sensitivity']} sensitivity +{cohorts['engineering_boundary']} boundaries") if cases else "не прочитан freeze"
        rows = self.manifest.get("methodology", [
            ["Политика", "fixed / failover", "host-monotonic", "параметры протокола сохранены"],
            ["Радиоконтроль", "geometry / LOS-only", "native ns-3", "одна физическая сцена"],
            ["Повторения", repeats, "независимый seed block", "вспомогательные когорты отдельно"],
            ["Горизонт", fmt(self.freeze.get("observation_horizon_source_sim_time_s")) + " с",
             "часы источника поз", "до просмотра эффекта политики"],
            ["Цели", "27 waypoint; барьеры", "автопилот time_boot_ms", "продвижение по радио, dwell3с"],
            ["Физический hold", "2 с", "source_sim_time", "отдельно от route dwell"],
            ["Timely", "250 мс", "command.created → UART", "offered остаётся знаменателем"],
            ["Канал", "n_LOS=2; n_NLOS=3", "2,412ГГц / 10dBm", "чувствительность2,7/3,3; без калибровки"],
            ["WP_SPD", "3 м/с", "полный MAVLink snapshot", "не жёсткая граница actual speed"],
        ])
        self.table(["Фактор", "Значение", "Шкала / источник", "Ограничение"], rows,
                   "Методика новой серии; незаполненные параметры не являются результатами.",
                   "New-study methodology; unresolved settings are not experimental results.")

    def legacy_table(self):
        path = PAPER / "data/six_cells.csv"
        self.inputs[str(path)] = sha(path)
        with path.open() as stream:
            rows = list(csv.DictReader(stream))
        self.table(["Отказ", "Политика", "Timely, %", "T_UART, с", "T_goal, с"],
            [[r["condition"], r["policy"], fmt(100*float(r["mean_timely_per_offered"])),
              fmt(r["mean_actual_control_s"]), fmt(r["mean_actual_goal_s"])] for r in rows],
            "Сохранённая вертикальная основная серия: пять запусков в каждой ячейке; Friis и заданный отказ.",
            "Retained vertical main study: five runs per cell; Friis and declared injected loss.")

    def results(self):
        descriptions = {
            "preparation": "Новая основная городская серия в данном черновике не представлена. Численные эффекты политики, интервалы и вывод об успешном восстановлении не заполнены. Это каркас методики и представления результатов.",
            "engineering_pilot": "Ниже представлены инженерные пилоты. Они не включены в оценку эффекта замороженной основной серии и не устанавливают воспроизводимый средний выигрыш failover.",
            "partial_series": "Ниже представлены доступные попытки ещё не завершённой серии. Неуспешные и неполные окна сохранены. Итоговый эффект по завершённому плану пока не заявляется.",
            "frozen_series": "Следующие таблицы относятся к основной части новой серии; вспомогательные когорты рассмотрены отдельно. Средние вычислены по полным измерительным окнам с сохранением статуса failed.",
        }
        self.para(descriptions[self.kind], "ispAnotation")
        if self.manifest.get("source_commit") and self.kind != "frozen_series":
            self.para("Версия новой реализации: " + str(self.manifest["source_commit"]) + ".")
        if self.manifest.get("probe_review"):
            probe = self.read_json(resolved(self.manifest["probe_review"], self.base))
            self.para("Компонентные проверки ниже относятся к неподвижным узлам и native Wi-Fi. Они проверяют реализацию геометрии/потерь и пакетного пути, но не заменяют пятибортовый полёт либо полевую калибровку.")
            self.table(["Проба", "LOS geometry", "Loss, dB", "Direct", "Via relay"],
                [[r["case"], r["geom_los"], fmt(r["loss_db"]), r["direct_received_unique"],
                  r["via_relay_received_unique"]] for r in probe["results"]],
                "Реальные компонентные пробы; уникальные sequence IDs проверены по PCAP.",
                "Measured component probes; unique sequence IDs checked against PCAP.")
        if self.study:
            self.study_results()
        if self.runs and self.manifest.get("per_run_table", not bool(self.study)):
            rows = []
            for label, group, run in self.runs:
                rate = run.get("rates", {}).get("all_five", {})
                endpoint = run.get("uart_after_NLOS", {}).get("all_five", {})
                time = fmt(endpoint.get("time_s")) if endpoint.get("status") == "observed" else {
                    "administratively_right_censored": ">" + fmt(endpoint.get("censor_time_s")),
                    "not_applicable_no_observed_NLOS_transition": "н/п: нет перехода",
                    "unverifiable_recording": "нет данных",
                    "interrupted_attempt": "прервано",
                }.get(endpoint.get("status"), "—")
                progress = list(run.get("mission_progress", {}).values())
                planned = sum(r["planned_waypoints"] for r in progress) if progress and all(
                    r.get("planned_waypoints") is not None for r in progress) else None
                done = sum(r["completed_waypoints"] for r in progress) if progress else None
                points = f"{done}/{planned}" if planned is not None and done is not None else "—"
                modes = "/".join(run.get("geometry", {}).get("selected_radio_modes", [])) or "—"
                timely = rate.get("timely_per_offered")
                rows.append([f"{label} ({group})", run.get("policy", "—") + "/" + modes,
                    run.get("recorded_status", "—"), fmt(100*timely) if timely is not None else "—",
                    time, points])
            self.table(["Опыт / группа", "Политика / radio", "Статус", "Timely, %", "После NLOS, с", "Точки / план"], rows,
                "Новые наблюдения по отдельным попыткам. Время после NLOS не называется восстановлением без подтверждённой деградации.",
                "New per-attempt observations. Post-NLOS latency is not called recovery without observed impairment.")
        for table in self.manifest.get("tables", []):
            path = resolved(table["csv"], self.base)
            self.inputs[str(path)] = sha(path)
            with path.open() as stream:
                source = list(csv.reader(stream))
            if not source:
                raise ValueError("empty numerical table")
            self.table(source[0], source[1:], table["caption_ru"], table["caption_en"])

    def study_results(self):
        study = self.study
        cell_rows = study_inputs.cells(study)
        effect_rows = study_inputs.contrasts(study, study_inputs.PRIMARY_METRICS + ("restricted_post_NLOS_UART_60s",))
        interaction_rows = study_inputs.interactions(study)
        analyses = {run["run_id"]: run for _, _, run in self.runs}
        stages = study_inputs.stage_counts(study, analyses)
        self.presentation_tables = dict(cells=cell_rows, contrasts=effect_rows,
                                       interactions=interaction_rows, stage_counts=stages)
        self.table(["Радио / политика", "Попытки / failed / полные", "Timely, % (n)", "Точки, % (n)"],
            [[r["radio_mode"] + "/" + r["policy"], f"{r['planned']}/{r['failed']}/{r['eligible']}",
              (fmt(100*r["timely_per_offered"]) + f" ({r['timely_per_offered_n']})") if r["timely_per_offered"] is not None else "—",
              (fmt(100*r["completed_waypoint_fraction"]) + f" ({r['completed_waypoint_fraction_n']})") if r["completed_waypoint_fraction"] is not None else "—"]
             for r in cell_rows if r["cohort"] == "main"],
            "Основная серия: средние полных окон; n — число наблюдений. Failed не означает отсутствие пригодного измерительного окна.",
            "Main study: complete-window means, with metric-specific n. A failed attempt may retain a valid measurement window.")
        def interval_text(value):
            return "—" if value is None else "[" + "; ".join(fmt(x) for x in value) + "]"
        self.table(["Радиорежим / показатель", "Полных / план пар", "Failover−fixed", "95% интервал", "Bounds пропусков"],
            [[r["radio_mode"] + "/" + r["metric_label"], f"{r['complete_pairs']}/{r['planned_pairs']}",
              fmt(r["mean"]), interval_text(r["interval_95"]), interval_text(r["missing_bounds"])]
             for r in effect_rows if r["cohort"] == "main"],
            "Основная новая серия: парные эффекты; интервалы разведочные. Bounds — границы чувствительности к отсутствующим исходам, не CI.",
            "New main study: paired effects with exploratory intervals. Missing-outcome bounds are not confidence intervals.")
        self.table(["Показатель", "Полных / план блоков", "Interaction, п.п.", "95% интервал", "Bounds пропусков"],
            [[r["metric_label"], f"{r['complete_blocks']}/{r['planned_blocks']}", fmt(r["mean"]),
              interval_text(r["interval_95"]), interval_text(r["missing_bounds"])] for r in interaction_rows],
            "Разность эффектов geometry−LOS-only требует четырёх случаев одного блока. Отсутствующее значение не заменено нулём.",
            "Geometry-minus-LOS-only policy interaction requires all four cases of a seed block. Missing is not zero.")
        self.table(["Радио / политика (n)", "Offered", "Создано CID", "App / UART", "Точки"],
            [[r["radio_mode"] + "/" + r["policy"] + f" ({r['complete_stage_runs']})",
              r["offered"], r["target_CIDs_created_in_window"],
              f"{r['follower_application_receive_CIDs']} / {r['UART_write_end_CIDs']}" if r["UART_write_end_CIDs"] is not None else "—",
              r["independently_verified_unique_inspection_goals"]] for r in stages if r["cohort"] == "main"],
            "Суммы по полным stage-экспортам основной серии. App receive — валидированный прикладной приём CID; это не PHY Rx/ACK. Цели — уникальные waypoint.",
            "Sums over complete stage exports of the main study. Application CID receipt is not PHY reception or ACK; goals are unique waypoints.")
        self.para("Счётчики CID относятся к целевым командам, созданным в окне, с проверенной идентичностью run/UAV/CID/epoch/relay, payload и цели. Все показанные UART-допуски прошли аудит полномочий. Native PHY-путь обеих передач и обратной телеметрии разобран отдельно для четырёх follower диагностической пары; это не покадровый аудит всех 56 попыток.")

    def figures(self):
        for entry in self.manifest.get("figures", []):
            if entry.get("measured") is not True or not self.runs:
                raise ValueError("urban figures require explicit measured:true and actual analysis inputs")
            picture = entry.get("path") or entry.get("paths", {}).get("png")
            if not picture:
                raise ValueError("figure requires path or paths.png")
            self.figure(resolved(picture, self.base), entry["caption_ru"], entry["caption_en"])

    def build(self):
        source = HERE / "src" / ("supplement.md" if self.args.supplement else "manuscript.md")
        self.inputs[str(source)] = sha(source)
        self.inputs[str(Path(__file__))] = sha(Path(__file__))
        for path in (HERE / "study_inputs.py", PAPER / "tools/build_docx.py",
                     PAPER / "template/official_russian_template.dotm"):
            self.inputs[str(path)] = sha(path)
        for line in source.read_text().splitlines():
            if not line.strip():
                continue
            prefixes = (("@title ", "ispHeader", ""), ("@title_en ", "ispHeader1", ""),
                ("@abstract ", "ispAnotation", "Аннотация. "), ("@abstract_en ", "ispAnotation", "Abstract. "),
                ("@keywords ", "ispAnotation", "Ключевые слова: "), ("@keywords_en ", "ispAnotation", "Keywords: "),
                ("## ", "ispSubHeader-3 level", ""), ("# ", "ispSubHeader-2 level", ""))
            match = next(((prefix, style, extra) for prefix, style, extra in prefixes if line.startswith(prefix)), None)
            if match:
                prefix, style, extra = match
                self.para(extra + line[len(prefix):], style)
            elif line == "@table methodology":
                self.methodology()
            elif line == "@legacy_table":
                self.legacy_table()
            elif line == "@legacy_figures":
                self.figure(PAPER / "figures/paired.png",
                    "Сохранённая вертикальная серия: парные эффекты и разведочные интервалы. Это отдельная когорта.",
                    "Retained vertical study: paired effects and exploratory intervals. This is a separate cohort.")
            elif line == "@results":
                self.results()
            elif line == "@figures":
                self.figures()
            elif line == "@references":
                self.para("Список литературы / References", "ispSubHeader-2 level")
                refs = HERE / "references.txt"
                self.inputs[str(refs)] = sha(refs)
                for ref in refs.read_text().splitlines():
                    if not self.args.supplement or int(ref.split("]")[0][1:]) <= 6:
                        self.para(ref, "ispLitList")
            elif line.startswith("@"):
                raise ValueError("unsupported manuscript directive: " + line)
            else:
                self.para(line)
        output = self.args.output.resolve()
        if output == PAPER / "build" or PAPER / "build" in output.parents:
            raise ValueError("previous paper/build is immutable")
        output.mkdir(parents=True, exist_ok=True)
        dest = output / "manuscript_anonymous.docx"
        self.doc.save(dest)
        with ZipFile(dest) as package:
            for name in package.namelist():
                if "vbaProject" in name or "attachedTemplate" in name:
                    raise ValueError("unexpected macro/template part")
                if name.endswith((".xml", ".rels")):
                    value = package.read(name)
                    if b"/home/" in value or b"file:///" in value or b"attachedTemplate" in value:
                        raise ValueError("local path/template link in anonymous document")
        result = dict(dataset_kind=self.kind, final_requested=self.args.final, input_sha256=self.inputs,
            analyses=len(self.runs), tables=self.table_number, figures=self.figure_number,
            docx_sha256=sha(dest), old_paper_build_written=False,
            note="document rendering validates presentation inputs, not scientific truth or model calibration")
        (output / "build_manifest.json").write_text(json.dumps(result, indent=2) + "\n")
        if hasattr(self, "presentation_tables"):
            (output / "presentation_tables.json").write_text(json.dumps(self.presentation_tables, indent=2) + "\n")
        print(dest)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--output", type=Path, default=HERE / "build")
    parser.add_argument("--supplement", action="store_true")
    parser.add_argument("--final", action="store_true")
    Revision(parser.parse_args()).build()


if __name__ == "__main__":
    main()
